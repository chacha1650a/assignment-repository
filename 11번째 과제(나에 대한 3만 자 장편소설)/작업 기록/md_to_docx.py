import re, sys
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

src = sys.argv[1]
dst = sys.argv[2]

text = open(src, encoding='utf-8').read()
lines = text.split('\n')

doc = Document()

style = doc.styles['Normal']
style.font.name = 'Malgun Gothic'
style.font.size = Pt(11)
rpr = style.element.get_or_add_rPr()
rFonts = rpr.find(qn('w:rFonts'))
if rFonts is None:
    from docx.oxml import OxmlElement
    rFonts = OxmlElement('w:rFonts')
    rpr.append(rFonts)
rFonts.set(qn('w:eastAsia'), 'Malgun Gothic')


def add_runs_with_bold(paragraph, text):
    # split on **bold** spans
    parts = re.split(r'(\*\*[^*]+\*\*)', text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**') and len(part) >= 4:
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        else:
            paragraph.add_run(part)


def is_table_line(line):
    return line.strip().startswith('|') and line.strip().endswith('|')


def parse_table(block_lines):
    rows = []
    for l in block_lines:
        l = l.strip()
        if not l.startswith('|'):
            continue
        cells = [c.strip() for c in l.strip('|').split('|')]
        rows.append(cells)
    # remove separator row (---|---|---)
    rows = [r for r in rows if not all(re.match(r'^:?-+:?$', c) for c in r)]
    return rows


def add_table(doc, rows):
    if not rows:
        return
    ncols = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = 'Light Grid Accent 1'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, row in enumerate(rows):
        for ci, cell_text in enumerate(row):
            if ci >= ncols:
                continue
            cell = table.cell(ri, ci)
            cell.text = ''
            p = cell.paragraphs[0]
            add_runs_with_bold(p, cell_text)
            if ri == 0:
                for run in p.runs:
                    run.bold = True


i = 0
n = len(lines)
first_title_done = False

while i < n:
    line = lines[i]
    stripped = line.strip()

    if stripped == '':
        i += 1
        continue

    if stripped == '---':
        i += 1
        continue

    if is_table_line(line):
        block = []
        while i < n and is_table_line(lines[i]):
            block.append(lines[i])
            i += 1
        rows = parse_table(block)
        add_table(doc, rows)
        doc.add_paragraph('')
        continue

    m = re.match(r'^(#{1,3})\s+(.*)$', line)
    if m:
        level = len(m.group(1))
        heading_text = m.group(2).strip()
        if not first_title_done and level == 1:
            p = doc.add_heading(heading_text, level=0)
            first_title_done = True
        else:
            hlevel = min(level, 3)
            doc.add_heading(heading_text, level=hlevel)
        i += 1
        continue

    if stripped == '[[ALEPH_BODY_START]]' or stripped == '[[ALEPH_BODY_END]]':
        p = doc.add_paragraph()
        p.add_run(stripped)
        i += 1
        continue

    if stripped.startswith('- '):
        p = doc.add_paragraph(style='List Bullet')
        add_runs_with_bold(p, stripped[2:].strip())
        i += 1
        continue

    if re.match(r'^\d+장\. ', stripped):
        p = doc.add_paragraph()
        add_runs_with_bold(p, stripped)
        i += 1
        continue

    # normal paragraph (possibly bold-only line like **김대훈**)
    p = doc.add_paragraph()
    add_runs_with_bold(p, stripped)
    p.paragraph_format.space_after = Pt(8)
    i += 1

doc.save(dst)
print('saved:', dst)
